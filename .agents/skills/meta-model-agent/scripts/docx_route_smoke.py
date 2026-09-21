from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
INIT = ROOT / "scripts" / "workspace_init.py"
EXPORT = ROOT / "scripts" / "docx_export.py"
BUILD_APPENDIX = ROOT / "scripts" / "build_code_appendix.py"
sys.path.insert(0, str(ROOT / "scripts"))
from gate_contracts import run_gate_check  # noqa: E402


def run(*args: str) -> None:
    proc = subprocess.run(
        [sys.executable, "-B", *args],
        text=True,
        capture_output=True,
        encoding="utf-8",
        errors="replace",
    )
    if proc.returncode:
        raise RuntimeError(proc.stdout + "\n" + proc.stderr)


def dense_sentence(competition: str) -> str:
    if competition == "mcm-icm":
        return "The model mechanism, objective, constraints, validated results, sensitivity analysis, interpretation, and limitations establish a robust recommendation. "
    return "模型机制、目标函数和约束已经推导，求解得到可靠结果；独立验证、灵敏度分析和结果解释说明方案稳定且具有现实意义。"


def valid_markdown(competition: str) -> str:
    sentence = dense_sentence(competition)
    if competition == "gmcm":
        return (
            "# 完整华为杯数学建模论文\n\n"
            "## 摘要\n本文针对容量约束下的资源配置问题，建立容量约束混合整数规划模型协调成本与服务覆盖。\n\n"
            "针对问题一，容量约束混合整数规划模型采用HiGHS分支定界算法求解，得到最优值10；约束回代验证表明结果可行。\n\n"
            "该模型结构清晰，具有可解释、稳健和可推广的优点，同时适用范围受容量参数假设限制。\n\n"
            "**关键词：** 混合整数线性规划；HiGHS；分支定界算法\n\n"
            "## 一、问题重述\n完整说明任务、输入、输出与评价准则。\n\n"
            "## 二、问题一模型的建立与求解\n\n"
            "### 模型建立与数学表达\n\n#### 变量、目标与约束\n建立容量约束混合整数规划模型。\n" + sentence * 135 +
            "\n\n### 求解方法与实现\n采用HiGHS分支定界算法求解，并记录参数与停止条件。\n" + sentence * 25 +
            "\n\n### 结果、检验与解释\n结果为10，约束回代、灵敏度分析和独立复核均表明结论稳定。\n" + sentence * 25 +
            "\n\n## 三、结论与模型评价\n逐问回答并说明适用边界、局限和推广条件。\n\n"
            "## 参考文献\n[1] 测试文献。\n\n"
            "<!-- CODE_APPENDIX_START -->\n## 附录 A：程序与复现说明\n<!-- CODE_APPENDIX_END -->"
        )
    if competition == "mcm-icm":
        return (
            "# Complete Modeling Paper\n\n**Control Number:** 1234567　　**Problem:** A\n\n"
            "## Summary\nThis study addresses capacity-constrained allocation and establishes a Capacity-Constrained Mixed-Integer Programming Model to balance cost and service coverage.\n\n"
            "For Problem 1, the Capacity-Constrained Mixed-Integer Programming Model is solved with HiGHS branch-and-bound, obtaining an optimal value of 10; constraint validation confirms feasibility.\n\n"
            "The model is interpretable, robust, and transferable, while its recommendations remain limited by the assumed capacity range.\n\n"
            "**Keywords:** Mixed-Integer Linear Programming; HiGHS; branch-and-bound\n\n"
            "## 1 Model Development, Results, and Validation\nThe Capacity-Constrained Mixed-Integer Programming Model is formulated as Mixed-Integer Linear Programming and solved by HiGHS branch-and-bound.\n" + sentence * 155 +
            "\n\n## 2 Discussion and Interpretation\n" + sentence * 40 +
            "\n\n## References\n[1] Test reference.\n\n<!-- CODE_APPENDIX_START -->\n## Appendix: Code and Reproducibility Notes\n<!-- CODE_APPENDIX_END -->"
            "\n\n## Report on Use of AI Tools\nNo AI-generated claim was accepted without human verification."
        )
    header = "**题号：** A　　**报名号：** 51001234　　**组别：** 本科\n\n" if competition == "51mcm" else ""
    return (
        "# 完整数学建模论文\n\n" + header +
        "## 摘要\n本文针对容量约束下的资源配置问题，建立容量约束混合整数规划模型协调成本与服务覆盖。\n\n"
        "针对问题一，容量约束混合整数规划模型采用HiGHS分支定界算法求解，得到最优值10；约束回代验证表明结果可行。\n\n"
        "该模型结构清晰，具有可解释、稳健和可推广的优点，同时适用范围受容量参数假设限制。\n\n"
        "**关键词：** 混合整数线性规划；HiGHS；分支定界算法\n\n"
        "## 问题一模型、结果与验证\n建立容量约束混合整数规划模型，采用HiGHS分支定界算法求解。\n" + sentence * 110 +
        "\n\n## 结论与解释\n" + sentence * 25 + "\n\n## 参考文献\n[1] 测试文献。"
        "\n\n<!-- CODE_APPENDIX_START -->\n## 附录：程序与复现说明\n<!-- CODE_APPENDIX_END -->"
    )


def prepare_model_context(workspace: Path, competition: str) -> None:
    (workspace / "问题分析.md").write_text(
        "本赛题共 1 个子问题。\n数据模式: none\n预处理: skipped\n本题无附件数据。\n",
        encoding="utf-8",
    )
    if competition == "mcm-icm":
        name = "Capacity-Constrained Mixed-Integer Linear Programming Model"
        family = "Mixed-Integer Linear Programming"
        algorithm = "HiGHS branch-and-bound algorithm"
        report = (
            "# Modeling Report\n\n## Problem 1\n"
            f"MODEL DEFINITION Q1 | ACADEMIC NAME: {name} | CANONICAL MODEL FAMILY: {family} | SOLVER ALGORITHM: {algorithm}\n"
            "MODEL STRUCTURE Q1 | DECISION/STATE VARIABLES: allocation and activation variables | OBJECTIVE/STATISTICAL RELATION: minimize total cost | CORE CONSTRAINTS/EQUATIONS: capacity and balance constraints | CUSTOM MECHANISM: problem-specific capacity limits\n"
            "PAPER EXPRESSION Q1 | DISPLAY NAME: Capacity-Constrained Mixed-Integer Programming Model | QUESTION ROLE: new_model | INHERITS FROM: none | CORE METHOD: HiGHS branch-and-bound\n"
            "Data mode: none; preprocessing: skipped.\n"
        )
    else:
        name = "考虑容量约束的混合整数线性规划模型"
        family = "混合整数线性规划"
        algorithm = "HiGHS分支定界算法"
        report = (
            "# 建模报告\n\n## 问题一\n"
            f"模型定义 Q1 | 正式名称: {name} | 标准模型族: {family} | 求解算法: {algorithm}\n"
            "模型结构 Q1 | 决策变量/状态量: 配置量与启用变量 | 目标函数/统计关系: 最小化总成本 | 核心约束/方程: 容量与平衡约束 | 定制机制: 题目容量上限\n"
            "论文表达 Q1 | 展示名称: 容量约束混合整数规划模型 | 问题角色: new_model | 继承模型: none | 核心方法: HiGHS分支定界算法\n"
            "数据模式: none\n预处理: skipped\n"
        )
    (workspace / "建模报告.md").write_text(report, encoding="utf-8")
    result = {
        "model_identity": {
            "Q1": {
                "academic_name": name,
                "canonical_model_family": family,
                "solver_algorithm": algorithm,
            }
        }
    }
    (workspace / "图表").mkdir(parents=True, exist_ok=True)
    (workspace / "图表" / "全部结果.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def prepare_code_appendix(workspace: Path) -> None:
    program = workspace / "程序" / "主程序.py"
    program.parent.mkdir(parents=True, exist_ok=True)
    program.write_text("def main():\n    return {'status': 'ok'}\n\nif __name__ == '__main__':\n    main()\n" * 8, encoding="utf-8")
    run(str(BUILD_APPENDIX), "--workspace", str(workspace), "--format", "markdown", "--insert-into", "论文/论文正文.md")


def ensure_page_measurement(workspace: Path, report: dict) -> dict:
    if report.get("page_count") is None:
        from pypdf import PdfWriter
        preview = workspace / "论文" / "docx_preview" / "smoke-preview.pdf"
        preview.parent.mkdir(parents=True, exist_ok=True)
        writer = PdfWriter()
        for _ in range(10):
            writer.add_blank_page(width=595, height=842)
        with preview.open("wb") as handle:
            writer.write(handle)
        report = dict(report)
        report.update({
            "page_count": 10,
            "body_page_count": 8,
            "abstract_page_count": 1,
            "preview_pdf": str(preview),
            "preview_sha256": hashlib.sha256(preview.read_bytes()).hexdigest(),
            "preview_mtime_ns": preview.stat().st_mtime_ns,
            "preview_engine": "smoke",
            "preview_error": None,
        })
        (workspace / "论文" / "docx_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def assert_competition_contract(base: Path, competition: str) -> dict:
    workspace = base / f"docx-{competition}"
    init_args = [str(INIT), "--workspace", str(workspace), "--competition", competition, "--output-format", "docx"]
    if competition == "gmcm":
        from docx import Document
        problem = base / "gmcm-problem.txt"
        official = base / "gmcm-official-template.docx"
        problem.write_text("GMCM smoke problem", encoding="utf-8")
        template = Document()
        template.add_paragraph("官方模板封面占位测试")
        template.save(official)
        init_args.extend(["--problem", str(problem), "--template", str(official), "--ai-disclosure", "off"])
    run(*init_args)
    prepare_model_context(workspace, competition)
    source = workspace / "论文" / "论文正文.md"
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_text(valid_markdown(competition), encoding="utf-8")
    prepare_code_appendix(workspace)
    good = run_gate_check(workspace, "MANUSCRIPT")
    assert good["passed"], good["issues"]
    original = source.read_text(encoding="utf-8")
    if competition == "51mcm":
        source.write_text(original.replace("51001234", "[报名号]"), encoding="utf-8")
        bad = run_gate_check(workspace, "MANUSCRIPT")
        assert not bad["passed"] and any("docx_51mcm_registration" in issue or "docx_placeholders" in issue for issue in bad["issues"]), bad
    elif competition == "mcm-icm":
        source.write_text(original.replace("1234567", "[CONTROL NUMBER]"), encoding="utf-8")
        missing_number = run_gate_check(workspace, "MANUSCRIPT")
        assert not missing_number["passed"] and any("docx_control_number" in issue or "docx_placeholders" in issue for issue in missing_number["issues"]), missing_number
        source.write_text(original + "\n中文混入。", encoding="utf-8")
        chinese = run_gate_check(workspace, "MANUSCRIPT")
        assert not chinese["passed"] and any("docx_english_only" in issue for issue in chinese["issues"]), chinese
    else:
        source.write_text(original + "\n\n##### 禁止的四级目录\n内容。", encoding="utf-8")
        deep = run_gate_check(workspace, "MANUSCRIPT")
        assert not deep["passed"] and any("gmcm_docx_toc_structure" in issue for issue in deep["issues"]), deep
    source.write_text(original, encoding="utf-8")
    result = {"competition": competition, "valid_gate": "pass", "negative_gate": "pass"}
    if competition == "gmcm":
        run(str(EXPORT), "--workspace", str(workspace))
        export_report = json.loads((workspace / "论文" / "docx_report.json").read_text(encoding="utf-8"))
        assert export_report["official_template_applied"] is True
        assert export_report["toc_included"] is True and export_report["toc_depth"] == 3
        with zipfile.ZipFile(workspace / "论文" / "数模论文.docx") as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        assert 'TOC \\o "1-3"' in document_xml
        assert all(style in document_xml for style in ("Heading1", "Heading2", "Heading3"))
        assert "官方模板封面占位测试" in document_xml
        from pypdf import PdfWriter
        manual_preview = workspace / "论文" / "manual-preview.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        with manual_preview.open("wb") as handle:
            writer.write(handle)
        run(str(EXPORT), "--workspace", str(workspace), "--preview-pdf", str(manual_preview))
        registered = json.loads((workspace / "论文" / "docx_report.json").read_text(encoding="utf-8"))
        assert registered["preview_engine"] == "manual"
        assert registered["preview_sha256"] == hashlib.sha256(manual_preview.read_bytes()).hexdigest()
        result["official_template"] = "pass"
        result["toc"] = "pass"
        result["manual_preview_registration"] = "pass"
    return result


def main() -> int:
    from PIL import Image
    base = Path(tempfile.mkdtemp(prefix="math-model-docx-route-"))
    workspace = base / "docx-route"
    run(str(INIT), "--workspace", str(workspace), "--competition", "cumcm", "--output-format", "docx")
    prepare_model_context(workspace, "cumcm")
    source = workspace / "论文" / "论文正文.md"
    source.parent.mkdir(parents=True, exist_ok=True)
    image_path = workspace / "论文" / "oversized_source.png"
    Image.new("RGB", (6000, 4000), "white").save(image_path)
    source.write_text(valid_markdown("cumcm").replace("\n\n## 结论与解释", "\n\n![超大原始图](oversized_source.png)\n\n## 结论与解释"), encoding="utf-8")
    prepare_code_appendix(workspace)
    run(str(EXPORT), "--workspace", str(workspace))
    report = json.loads((workspace / "论文" / "docx_report.json").read_text(encoding="utf-8"))
    report = ensure_page_measurement(workspace, report)
    assert (workspace / "论文" / "数模论文.docx").stat().st_size >= 15000
    assert report["effective_body_units"] >= 3500
    assert report["images"]
    assert all(item["width_cm"] <= report["usable_width_cm"] for item in report["images"])
    manuscript_gate = run_gate_check(workspace, "MANUSCRIPT")
    assurance_gate = run_gate_check(workspace, "ASSURANCE")
    assert manuscript_gate["passed"], manuscript_gate["issues"]
    assert assurance_gate["passed"], assurance_gate["issues"]
    tampered = dict(report)
    tampered["images"] = [dict(report["images"][0], width_cm=report["usable_width_cm"] + 5)]
    (workspace / "论文" / "docx_report.json").write_text(json.dumps(tampered, ensure_ascii=False, indent=2), encoding="utf-8")
    rejected = run_gate_check(workspace, "ASSURANCE")
    assert not rejected["passed"] and any("docx_image_size_contract" in issue for issue in rejected["issues"])
    (workspace / "论文" / "docx_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    source.write_text(source.read_text(encoding="utf-8") + "\n新增但尚未导出的解释。", encoding="utf-8")
    stale = run_gate_check(workspace, "ASSURANCE")
    assert not stale["passed"] and any("docx_source_freshness" in issue or "docx_output_freshness" in issue for issue in stale["issues"]), stale
    run(str(EXPORT), "--workspace", str(workspace))
    report = json.loads((workspace / "论文" / "docx_report.json").read_text(encoding="utf-8"))
    report = ensure_page_measurement(workspace, report)
    assert run_gate_check(workspace, "ASSURANCE")["passed"]
    report["competition_contracts"] = [
        assert_competition_contract(base, "51mcm"),
        assert_competition_contract(base, "mcm-icm"),
        assert_competition_contract(base, "gmcm"),
    ]
    print(json.dumps(report, ensure_ascii=False, indent=2))
    shutil.rmtree(base, ignore_errors=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
