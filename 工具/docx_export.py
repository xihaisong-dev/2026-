from __future__ import annotations

import argparse
import base64
import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path

def image_dimensions(path: Path) -> tuple[int, int]:
    try:
        from PIL import Image
        with Image.open(path) as image:
            return image.size
    except Exception:
        return (1600, 900)


def effective_units(text: str) -> int:
    text = re.sub(r"```.*?```", " ", text, flags=re.DOTALL)
    text = re.sub(r"!\[[^]]*\]\([^)]+\)", " ", text)
    text = re.sub(r"[#*_>`|~-]", " ", text)
    return len(re.findall(r"[\u4e00-\u9fff]", text)) + len(re.findall(r"\b[A-Za-z][A-Za-z'-]*\b", text))


def uploaded_word_templates(workspace: Path, state: dict) -> list[Path]:
    records = (state.get("problem_intake") or {}).get("template_files") or []
    candidates = [workspace / str(item.get("path", "")) for item in records]
    root = workspace / "题目" / "官方论文模板"
    if root.exists():
        candidates.extend(path for path in root.rglob("*") if path.is_file())
    unique: dict[Path, Path] = {}
    for path in candidates:
        resolved = path.resolve()
        if resolved.exists() and resolved.suffix.lower() in {".doc", ".docx"}:
            unique[resolved] = resolved
    return sorted(unique.values(), key=lambda path: (path.suffix.lower() != ".docx", path.name.casefold()))


def convert_with_word(source: Path, destination: Path, output_kind: str) -> str | None:
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        return "Windows PowerShell is unavailable"
    source_literal = str(source.resolve()).replace("'", "''")
    destination_literal = str(destination.resolve()).replace("'", "''")
    operation = (
        "$document.ExportAsFixedFormat($target, 17)"
        if output_kind == "pdf"
        else "$document.SaveAs2($target, 16)"
    )
    script = f"""
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
$source = '{source_literal}'
$target = '{destination_literal}'
try {{
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($source, $false, $true)
    {operation}
}} finally {{
    if ($null -ne $document) {{ $document.Close(0) }}
    if ($null -ne $word) {{ $word.Quit() }}
}}
"""
    encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        destination.unlink()
    proc = subprocess.run(
        [powershell, "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
        capture_output=True,
        text=True,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    if proc.returncode or not destination.exists():
        return (proc.stderr or proc.stdout or "Microsoft Word conversion did not create the requested file").strip()
    return None


def resolve_official_docx(workspace: Path, state: dict, soffice: str | None) -> tuple[Path, Path, str]:
    candidates = uploaded_word_templates(workspace, state)
    if not candidates:
        raise SystemExit("GMCM DOCX export requires the current official .doc or .docx paper template uploaded with the problem.")
    source = candidates[0]
    if source.suffix.lower() == ".docx":
        return source, source, "not_needed"
    cache = workspace / "论文" / "docx_template_cache"
    cache.mkdir(parents=True, exist_ok=True)
    converted = cache / f"{source.stem}.docx"
    conversion_errors: list[str] = []
    if converted.exists():
        converted.unlink()
    if soffice:
        proc = subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", str(cache), str(source)],
            capture_output=True,
            text=True,
        )
        if not proc.returncode and converted.exists():
            return source, converted, "libreoffice"
        conversion_errors.append(f"LibreOffice: {proc.stderr or proc.stdout}")
    word_error = convert_with_word(source, converted, "docx")
    if word_error is None:
        return source, converted, "microsoft_word"
    conversion_errors.append(f"Microsoft Word: {word_error}")
    raise SystemExit(
        "The uploaded GMCM template is legacy .doc and could not be converted. "
        "Install LibreOffice, make Microsoft Word available on Windows, or upload the official .docx version. "
        + " | ".join(conversion_errors)
    )


def add_toc_field(document, OxmlElement, qn, WD_ALIGN_PARAGRAPH) -> None:
    title = document.add_paragraph("目录")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title.runs:
        title.runs[0].bold = True
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开文档后更新目录字段"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, placeholder, end):
        run._r.append(node)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)


def markdown_heading_level(markers: str, competition: str) -> int | None:
    depth = len(markers)
    if competition == "gmcm":
        if depth == 1:
            return None
        if depth > 4:
            raise SystemExit("GMCM permits at most three directory levels; Markdown headings deeper than #### are forbidden.")
        return depth - 1
    return min(depth, 3)


def ensure_heading_styles(document, WD_STYLE_TYPE, OxmlElement, qn, Pt) -> None:
    styles = document.styles
    normal = styles["Normal"]
    specifications = [
        ("Title", 18, True, None),
        ("Heading 1", 16, True, 0),
        ("Heading 2", 14, True, 1),
        ("Heading 3", 12, True, 2),
    ]
    for name, size, bold, outline_level in specifications:
        try:
            style = styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
        style.font.size = Pt(size)
        style.font.bold = bold
        if outline_level is not None:
            ppr = style.element.get_or_add_pPr()
            existing = ppr.find(qn("w:outlineLvl"))
            if existing is not None:
                ppr.remove(existing)
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), str(outline_level))
            ppr.append(outline)


def measure_preview(report: dict, preview: Path, competition: str) -> dict:
    report["preview_pdf"] = str(preview)
    report["preview_sha256"] = hashlib.sha256(preview.read_bytes()).hexdigest()
    report["preview_mtime_ns"] = preview.stat().st_mtime_ns
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(preview))
        report["page_count"] = len(reader.pages)
        page_texts = [(page.extract_text() or "") for page in reader.pages]
        if competition == "cumcm":
            body_start = next((index for index, value in enumerate(page_texts, 1) if "问题重述" in value), None)
            appendix_start = next((index for index, value in enumerate(page_texts, 1) if "附录" in value and "程序" in value), None)
            if body_start is not None:
                report["abstract_page_count"] = max(body_start - 1, 0)
            if body_start is not None and appendix_start is not None and appendix_start >= body_start:
                report["body_page_count"] = appendix_start - body_start + 1
        if competition == "gmcm":
            abstract_start = next((index for index, value in enumerate(page_texts, 1) if "摘要" in value and "关键词" in value), None)
            toc_start = next((index for index, value in enumerate(page_texts, 1) if "目录" in value), None)
            body_start = next((index for index, value in enumerate(page_texts, 1) if "问题重述" in value or "问题分析" in value), None)
            appendix_start = next((index for index, value in enumerate(page_texts, 1) if "附录 A" in value or "附录A" in value), None)
            if abstract_start is not None and toc_start is not None and toc_start >= abstract_start:
                report["abstract_page_count"] = max(toc_start - abstract_start, 1)
            if body_start is not None:
                body_end = appendix_start if appendix_start is not None and appendix_start >= body_start else len(page_texts)
                report["body_page_count"] = body_end - body_start + 1
        report["preview_error"] = None
    except Exception as exc:
        report["preview_error"] = str(exc)
    return report


def register_preview(workspace: Path, preview: Path) -> dict:
    report_path = workspace / "论文" / "docx_report.json"
    source = workspace / "论文" / "论文正文.md"
    output = workspace / "论文" / "数模论文.docx"
    if not report_path.exists() or not source.exists() or not output.exists():
        raise SystemExit("Export the current DOCX before registering a manual PDF preview.")
    preview = preview.resolve()
    if not preview.is_file():
        raise SystemExit(f"manual PDF preview does not exist: {preview}")
    report = json.loads(report_path.read_text(encoding="utf-8"))
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    if report.get("source_sha256") != source_hash:
        raise SystemExit("The Markdown source changed after DOCX export; export the DOCX again before registering its PDF preview.")
    report["preview_engine"] = "manual"
    measure_preview(report, preview, str(report.get("competition", "cumcm")))
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return report


def export(workspace: Path) -> dict:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt
    except Exception as exc:
        raise SystemExit("python-docx is required: pip install python-docx") from exc
    state = json.loads((workspace / "状态" / "工作流状态.json").read_text(encoding="utf-8"))
    competition = state.get("competition", "cumcm")
    source = workspace / "论文" / "论文正文.md"
    output = workspace / "论文" / "数模论文.docx"
    if not source.exists():
        raise SystemExit(f"missing source: {source}")
    text = source.read_text(encoding="utf-8")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    official_template: Path | None = None
    official_template_source: Path | None = None
    template_conversion_engine: str | None = None
    if competition == "gmcm":
        official_template_source, official_template, template_conversion_engine = resolve_official_docx(workspace, state, soffice)
        try:
            document = Document(str(official_template))
        except Exception as exc:
            raise SystemExit(f"The uploaded official GMCM Word template is not a readable DOCX file: {official_template}") from exc
        document.add_page_break()
    else:
        document = Document()
    section = document.sections[0]
    if competition == "mcm-icm":
        section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    elif competition != "gmcm":
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    if competition != "gmcm":
        section.left_margin = section.right_margin = Cm(2.54)
        section.top_margin = section.bottom_margin = Cm(2.54)
    usable_width_cm = (section.page_width - section.left_margin - section.right_margin) / 360000
    usable_height_cm = (section.page_height - section.top_margin - section.bottom_margin) / 360000
    normal = document.styles["Normal"].font
    if competition != "gmcm" or not normal.name:
        normal.name = "Times New Roman" if competition == "mcm-icm" else "宋体"
    if competition != "gmcm" or normal.size is None:
        normal.size = Pt(12 if competition in {"mcm-icm", "gmcm"} else 10.5)
    ensure_heading_styles(document, WD_STYLE_TYPE, OxmlElement, qn, Pt)
    images: list[dict] = []
    seen_abstract = False
    toc_included = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line or re.fullmatch(r"<!--.*-->", line):
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            image_path = Path(image_match.group(2))
            if not image_path.is_absolute():
                image_path = (source.parent / image_path).resolve()
            if not image_path.exists():
                raise SystemExit(f"missing image: {image_path}")
            px_w, px_h = image_dimensions(image_path)
            ratio = px_h / max(px_w, 1)
            width_cm = min(15.0, float(usable_width_cm))
            height_cm = width_cm * ratio
            max_height_cm = min(20.0, float(usable_height_cm) * 0.72)
            if height_cm > max_height_cm:
                height_cm = max_height_cm
                width_cm = height_cm / max(ratio, 0.01)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm), height=Cm(height_cm))
            if image_match.group(1):
                caption = document.add_paragraph(image_match.group(1))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            images.append({"path": str(image_path), "width_cm": round(width_cm, 3), "height_cm": round(height_cm, 3)})
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            heading_text = heading.group(2).strip()
            if competition == "gmcm" and seen_abstract and not toc_included and "摘要" not in heading_text:
                document.add_page_break()
                add_toc_field(document, OxmlElement, qn, WD_ALIGN_PARAGRAPH)
                document.add_page_break()
                toc_included = True
            level = markdown_heading_level(heading.group(1), competition)
            if level is None:
                document.add_paragraph(heading_text, style="Title")
            else:
                document.add_heading(heading_text, level=level)
            if competition == "gmcm" and "摘要" in heading_text:
                seen_abstract = True
        else:
            document.add_paragraph(line)
    if competition == "gmcm" and not toc_included:
        raise SystemExit("GMCM DOCX source must contain an abstract followed by body headings so a three-level table of contents can be generated.")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    report = {
        "output_format": "docx",
        "competition": competition,
        "source": str(source),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_mtime_ns": source.stat().st_mtime_ns,
        "output": str(output),
        "effective_body_units": effective_units(text),
        "usable_width_cm": round(float(usable_width_cm), 3),
        "usable_height_cm": round(float(usable_height_cm), 3),
        "images": images,
        "page_count": None,
        "body_page_count": None,
        "abstract_page_count": None,
        "preview_pdf": None,
        "preview_sha256": None,
        "preview_mtime_ns": None,
        "toc_included": toc_included,
        "toc_depth": 3 if competition == "gmcm" else None,
        "official_template": str(official_template) if official_template else None,
        "official_template_source": str(official_template_source) if official_template_source else None,
        "official_template_sha256": hashlib.sha256(official_template_source.read_bytes()).hexdigest() if official_template_source else None,
        "converted_template_sha256": hashlib.sha256(official_template.read_bytes()).hexdigest() if official_template and official_template != official_template_source else None,
        "official_template_conversion_engine": template_conversion_engine,
        "official_template_applied": official_template is not None,
        "preview_engine": None,
        "preview_error": None,
    }
    if soffice:
        preview_dir = workspace / "论文" / "docx_preview"
        preview_dir.mkdir(parents=True, exist_ok=True)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(preview_dir), str(output)], capture_output=True, text=True)
        preview = preview_dir / "数模论文.pdf"
        if preview.exists():
            report["preview_engine"] = "libreoffice"
    else:
        preview = None
        report["preview_error"] = "automatic PDF preview requires LibreOffice; alternatively save a PDF manually and register it with --preview-pdf"
    if preview is not None and preview.exists():
        measure_preview(report, preview, competition)
    (workspace / "论文" / "docx_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Export Markdown manuscript to bounded DOCX.")
    parser.add_argument("--workspace", default=".")
    parser.add_argument("--preview-pdf", help="Register a manually exported PDF preview without regenerating the DOCX.")
    args = parser.parse_args()
    workspace = Path(args.workspace).resolve()
    if args.preview_pdf:
        register_preview(workspace, Path(args.preview_pdf))
    else:
        export(workspace)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
