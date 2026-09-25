#!/usr/bin/env python3
"""Build the complete Chinese paper from real, independently validated results.

Requires Pandoc on PATH and python-docx. Final mode refuses incomplete results,
unvalidated records, missing figures, duplicate case/core/scenario keys or unresolved
template tokens. --allow-partial is for internal layout drafts only.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
import gzip
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
BODY_FONT='Noto Serif CJK SC'
HEAD_FONT='Noto Sans CJK SC'
MI=2**20


def jread(path, default=None):
    if not Path(path).exists():return {} if default is None else default
    opener=gzip.open if str(path).endswith('.gz') else open
    with opener(path,'rt',encoding='utf-8') as f:return json.load(f)


def number(x,d=3):
    return '—' if pd.isna(x) else f'{x:.{d}f}'


def integer(x):
    if pd.isna(x):return '—'
    if float(x)!=int(x):raise ValueError(f'Expected exact integer count, got {x}')
    return str(int(x))


def table(headers,rows,caption):
    def esc(s):return str(s).replace('|','／').replace('\n',' ')
    return '\n\n'+caption+'\n\n'+'| '+' | '.join(map(esc,headers))+' |\n|'+ '|'.join([':---:']*len(headers))+'|\n'+''.join('| '+' | '.join(map(esc,row))+' |\n' for row in rows)+'\n'


def figure(name,caption,args):
    p=args.figures/(name+'.png')
    if not p.exists():
        if args.allow_partial:return f'\n\n【内部排版草稿：待生成 {name}】\n\n'
        raise FileNotFoundError(p)
    # Relative paths keep the packaged manuscript portable.
    rel=os.path.relpath(p,ROOT).replace(os.sep,'/')
    return f'\n\n![]({rel}){{width=15.5cm}}\n\n{caption}\n\n'


def validate_input(data,args):
    keys=['case','problem','cores']
    if data.duplicated(keys).any():raise ValueError('Duplicate case/problem/core records')
    expected={(f'case_{i:03d}',p,n) for i in range(1,101) for p in (1,2,3) for n in range(1,6)}
    actual=set(map(tuple,data[keys].itertuples(index=False,name=None)))
    if not actual<=expected:raise ValueError('Unexpected case, problem or core number')
    if not args.allow_partial and actual!=expected:raise ValueError(f'Final paper requires all 1500 records; have {len(actual)}; missing {len(expected-actual)}')
    if data.makespan.isna().any() or (data.makespan<=0).any():raise ValueError('Invalid makespan')
    if not np.allclose(data.speedup,data.baseline_makespan/data.makespan):raise ValueError('Speedup is inconsistent with official makespan')
    report=jread(args.results.parent/'validation_report.json')
    if not args.allow_partial:
        if report.get('status')!='PASS' or report.get('validated_results')!=1500 or report.get('missing_results'):
            raise ValueError('Full independent validation report is required before the final paper')
    paired=data[data.problem.isin([2,3])].pivot(index=['case','cores'],columns='problem',values='makespan').dropna()
    cached=data[data.problem==3].set_index(['case','cores'])
    if len(paired) and not np.allclose(paired[2]/paired[3],cached.loc[paired.index,'cache_speedup_vs_same_plan']):
        raise ValueError('P3 paired speedup inconsistent with P2/P3 makespans')


def make_tokens(data,args):
    toks={}; g=data.groupby(['problem','cores']);stats=g.agg(mean_speedup=('speedup','mean'),median_speedup=('speedup','median'),count=('case','count'))
    def subset(p,n=5):return data[(data.problem==p)&(data.cores==n)]
    def m(p,n,col='speedup'):
        s=subset(p,n);return s[col].mean() if len(s) else np.nan
    def seq(p,col='speedup',start=2):return '、'.join(number(m(p,n,col)) for n in range(start,6))
    toks.update(A_SPEEDUPS=seq(1),B_SPEEDUPS=seq(2),CACHE_SPEEDUPS=seq(3,'cache_speedup_vs_same_plan',1),
                A_SPEEDUP5=number(m(1,5)),B_SPEEDUP5=number(m(2,5)),
                CACHE_SPEEDUP5=number(m(3,5,'cache_speedup_vs_same_plan')),
                CACHE_HIT5=number(m(3,5,'cache_hit_rate')*100,2)+'%',RESULT_COUNT=str(len(data)))
    q=subset(1).merge(subset(2),on=['case','cores'],suffixes=('_A','_B'))
    ratio=q.makespan_A/q.makespan_B
    faster=int((q.makespan_B<q.makespan_A).sum());slower=int((q.makespan_B>q.makespan_A).sum());equal=int((q.makespan_B==q.makespan_A).sum())
    toks['B_OVER_A']=number(ratio.mean())
    toks['B_COMPARE_SENTENCE']=f'其中 {faster} 个用例在场景 B 下更快、{equal} 个相同、{slower} 个更慢'
    f=pd.read_csv(args.results.parent/'dataset_features.csv')
    features=[('张量节点数','n_tensors','个',1),('全部操作节点数','n_ops','个',1),('计算操作节点数','n_compute_ops','个',1),('原始有向边数','n_edges','条',1),('计算弱连通分量','n_components','个',1),('原始 DDR 搬运量','original_ddr_bytes','MiB',MI)]
    rows=[]
    for label,col,unit,scale in features:
        x=f[col]/scale;d=3 if scale!=1 else 1
        rows.append([label,unit,integer(x.min()) if scale==1 else number(x.min(),d),
                     number(x.median(),d),integer(x.max()) if scale==1 else number(x.max(),d)])
    toks['DATASET_TABLE']=table(['数据特征','单位','最小值','中位数','最大值'],rows,'表1 100个正式计算图的结构与搬运规模')
    toks['ALGORITHM_TABLE']=table(['候选类别','主要操作','适用目的'],[
        ['分量轮转','完整分量依次分配到核心','形成确定性的朴素对照'],
        ['分量负载均衡','按 LPT 或 Pipe 负载选择核心','降低粗粒度负载差异'],
        ['分量分批','场景 A 中按完整分量组织多个 Task','比较局部缓存调度与启动等待'],
        ['链与深度窗口','保留链局部性或按拓扑层次细分','缓解巨大分量限制并行的情况'],
        ['输入复用顺序','场景 B 中按共享输入身份排序','改变同核优先级与张量存活期'],
        ['整图回退候选','保留单个整图子图并填充空核心','提供稳定可行的比较方案']], '表2 结构候选及其作用')
    env_original=jread(args.results.parent/'execution_environment.json')
    env_certified=jread(args.results.parent/'certified_execution_environment.json')
    if (env_original.get('budget') is not None and env_certified.get('budget') is not None
            and env_original['budget']!=env_certified['budget']):
        raise ValueError('Original and certified execution environments disagree on candidate budget')
    env={**env_original,**env_certified};budget=env.get('budget','standard')
    toks['PYTHON_VERSION']=env.get('python',sys.version).split()[0]
    budget_detail=''
    if budget in ('standard','quick'):
        budget_detail='对超过 10,000 个计算操作节点的大图，不生成 LPT 与输入复用次序候选，保留分量轮转、Pipe 均衡以及满足各自结构触发条件的链切分和场景 A 分量分批候选；大图的深度窗口仅保留场景 A 的 16 层窗口和场景 B 的 4 层窗口。'
        budget_detail+=('未超过该规模阈值的图在触发细分时使用 4、8、16 层窗口。' if budget=='standard' else '未超过该规模阈值的图在触发细分时使用 4 层窗口；quick 模式最终截取去重后构造序列的前 4 个候选。')
    elif budget=='extended':
        budget_detail='extended 模式不启用大图候选缩减，在结构触发时进一步保留 32、64 层窗口和较强通信局部性候选。'
    toks['SOLVER_PROTOCOL']=f'本次采用 {budget} 候选预算。'+budget_detail+'候选数量随结构触发与方案去重而变化，并非每个图都评估相同数量的候选。所有参数保持固定，候选生成不读取用例编号进行人工分组；官方评估失败的候选记录错误原因后剔除。整图单核方案提供安全回退，多核方案可以保留空闲核心。相同方案按完整映射与核心顺序去重，避免重复评估。最终方案与每个候选的指标均随代码包保存。'
    toks['SPEEDUP_TABLE']=table(['核心数','场景 A 平均','场景 A 中位数','场景 B 平均','场景 B 中位数'],[[n,number(m(1,n)),number(subset(1,n).speedup.median()),number(m(2,n)),number(subset(2,n).speedup.median())] for n in range(1,6)],'表3 两种场景相对整图单核的加速比')
    a=subset(1);b=subset(2);c=subset(3)
    toks['Q1_RESULTS_TEXT']=f'在 5 核配置下，场景 A 的平均加速比为 {number(a.speedup.mean())}，中位数为 {number(a.speedup.median())}，最小值与最大值分别为 {number(a.speedup.min())} 和 {number(a.speedup.max())}。{int((a.speedup>1+1e-10).sum())} 个用例相对单核获得加速，{int((a.speedup>5+1e-10).sum())} 个用例的加速比超过 5。四分位区间为 [{number(a.speedup.quantile(.25))}, {number(a.speedup.quantile(.75))}]，说明平均值之外仍存在明显实例差异。'
    toks['Q2_RESULTS_TEXT']=f'场景 B 在 5 核下的平均加速比为 {number(b.speedup.mean())}，中位数为 {number(b.speedup.median())}，四分位区间为 [{number(b.speedup.quantile(.25))}, {number(b.speedup.quantile(.75))}]。相同用例与核数下，场景 A 与场景 B 的逐例 Makespan 比值平均为 {number(ratio.mean())}；场景 B 更快、相同和更慢的用例数分别为 {faster}、{equal} 和 {slower}。这个比较同时包含场景语义与各自方案选择的影响，不意味着场景 B 对每个图都必然占优。'
    toks['MOVEMENT_TABLE']=table(['核心数','A额外量/MiB','B额外量/MiB','A spill/MiB','B spill/MiB'],[[n,number(m(1,n,'added_copy_bytes')/MI),number(m(2,n,'added_copy_bytes')/MI),number(m(1,n,'spill_added_copy_bytes')/MI),number(m(2,n,'spill_added_copy_bytes')/MI)] for n in range(1,6)],'表4 额外搬运与缓存换入换出量的逐例均值')
    delta=(q.added_copy_bytes_A-q.added_copy_bytes_B).mean()/MI
    toks['MOVEMENT_RESULTS_TEXT']=f'5 核下场景 A 与场景 B 的平均额外搬运量分别为 {number(a.added_copy_bytes.mean()/MI)} MiB 和 {number(b.added_copy_bytes.mean()/MI)} MiB，二者之差（A−B）为 {number(delta)} MiB。对应平均 spill 搬运量分别为 {number(a.spill_added_copy_bytes.mean()/MI)} MiB 和 {number(b.spill_added_copy_bytes.mean()/MI)} MiB。图中同时列出切分搬运与 spill，能够区分边界通信减少和私有缓存压力变化；这些量与 Makespan 的关系还受计算重叠及搬运所在关键路径影响。'
    cache_rows=[]
    for n in range(1,6):
        s=subset(3,n);den=s.cache_hit_bytes.sum()+s.cache_miss_bytes.sum();pooled=s.cache_hit_bytes.sum()/den if den else 0
        cache_rows.append([n,number(m(2,n)),number(m(3,n)),number(m(3,n,'cache_speedup_vs_same_plan')),number(m(3,n,'cache_hit_rate')*100,2),number(pooled*100,2)])
    toks['CACHE_TABLE']=table(['核心数','无L2平均加速','有L2平均加速','配对相对加速','平均命中/%','合并字节命中/%'],cache_rows,'表5 只读Cache的同方案成对比较')
    cache_ratio=c.cache_speedup_vs_same_plan
    toks['Q3_RESULTS_TEXT']=f'5 核只读 Cache 相对无 L2 的平均加速比为 {number(cache_ratio.mean())}，中位数为 {number(cache_ratio.median())}，范围为 {number(cache_ratio.min())}～{number(cache_ratio.max())}。逐例平均字节命中率为 {number(c.cache_hit_rate.mean()*100,2)}%。有 L2 后 Makespan 更小、相同和更大的用例数分别为 {int((cache_ratio>1+1e-10).sum())}、{int(np.isclose(cache_ratio,1,rtol=0,atol=1e-10).sum())} 和 {int((cache_ratio<1-1e-10).sum())}。结果保留性能相同或下降的用例，没有只选取受益实例计算均值。'
    representative=pd.concat([c.nsmallest(3,'cache_speedup_vs_same_plan'),c.nlargest(3,'cache_speedup_vs_same_plan')]).drop_duplicates('case')
    toks['CACHE_CASE_TABLE']=table(['用例','无L2周期','有L2周期','相对加速','字节命中/%'],[[r.case,integer(b.set_index('case').loc[r.case,'makespan']),integer(r.makespan),number(r.cache_speedup_vs_same_plan),number(r.cache_hit_rate*100,2)] for r in representative.itertuples()], '表6 5核Cache收益分布两端的代表用例')
    corr=c.cache_hit_rate.rank().corr(cache_ratio.rank()) if len(c)>1 else np.nan
    toks['CACHE_DISCUSSION']=f'图中将命中率与配对加速比逐点对应，避免仅凭平均命中率判断整体性能。5 核下二者的 Spearman 秩相关系数为 {number(corr)}；这一描述性关系不构成因果识别。命中的 COPY_IN 是否处于关键路径、命中释放的 DDR 带宽能否缩短其他核心等待，以及 FIFO 访问次序的变化，都会影响最终收益。'
    toks['CONFIG_TABLE']=table(['资源或参数','固定值'],[['L1 / UB 容量','524288 / 131072 bytes'],['共享 DDR 总带宽','60 bytes/cycle'],['场景 A 同核 / 跨核等待','100 / 1000 cycles'],['场景 B 跨核 COPY 同步等待','500 cycles'],['共享只读 L2 容量','1048576 bytes'],['共享只读 L2 读带宽','250 bytes/cycle'],['核数与用例','1～5 核，100 个正式计算图']], '表7 统一硬件与实验配置')
    candidates_path=args.results.parent/'candidates.csv'
    cand=pd.read_csv(candidates_path,low_memory=False) if candidates_path.exists() else pd.DataFrame()
    rr=cand[(cand.method=='component_round_robin')&(cand.status=='ok')] if len(cand) else pd.DataFrame()
    if len(rr):
        rr=rr[['case','problem','cores','makespan']].rename(columns={'makespan':'baseline_rr'})
        z=data[data.problem.isin([1,2])&(data.cores>1)].merge(rr,on=['case','problem','cores'],validate='one_to_one');z['portfolio_gain']=z.baseline_rr/z.makespan
    else:z=pd.DataFrame()
    rows=[]
    if len(z):
        for (problem,n),s in z.groupby(['problem','cores']):rows.append(['A' if problem==1 else 'B',int(n),len(s),number(s.portfolio_gain.mean()),number(s.portfolio_gain.median()),int((s.portfolio_gain>1+1e-10).sum())])
    toks['ABLATION_TEXT']='采用保持完整弱连通分量的轮转分核作为统一朴素对照。下表在相同图、相同核数与相同场景上计算“轮转 Makespan / 最终方案 Makespan”，仅统计两者均成功评估的配对记录。比值大于 1 表示候选组合与官方评估选优带来额外收益。候选去重可能使部分策略不单独出现，因此这里不将缺失候选填成虚构测量。'
    toks['ABLATION_TABLE']=table(['场景','核数','配对数','平均相对加速','中位相对加速','严格改善数'],rows,'表8 最终候选组合相对分量轮转的收益')
    report=jread(args.results.parent/'validation_report.json');invalid=int(data.invalid_candidate_count.sum())
    toks['VALIDATION_TEXT']=f'独立校验已核对 {len(data)} 条正式结果，覆盖节点映射、子图调度、Makespan、逻辑搬运恒等式、容量峰值及问题二与问题三的方案一致性。正式结果的重复键数量为 0，结果校验状态为 {report.get("status","未完成")}。候选搜索记录了 {invalid} 次不可行候选，这些候选均被剔除，未进入统计。全部正式结果均保留原始官方 JSON，便于按用例复核。'
    # New certified-runner records explicitly separate winner verification.
    # Their search time already includes official preparations, the exact
    # effective-precedence certificate and candidate simulation. Old records have no explicit
    # final field and need the earlier reused-baseline audit-log correction.
    runtime_data=data.copy()
    runtime_data['total_evaluation_seconds']=runtime_data.search_evaluation_seconds
    explicit_final=pd.to_numeric(data.get('final_verification_seconds',pd.Series(np.nan,index=data.index)),errors='coerce')
    has_explicit_final=explicit_final.notna()
    runtime_data.loc[has_explicit_final,'total_evaluation_seconds']+=explicit_final[has_explicit_final]
    if 'total_evaluation_seconds' in data:
        saved_total=pd.to_numeric(data.total_evaluation_seconds,errors='coerce')
        check=has_explicit_final&saved_total.notna()
        if not np.allclose(runtime_data.loc[check,'total_evaluation_seconds'],saved_total[check]):
            raise ValueError('Explicit total evaluation time disagrees with search + final verification')
    for idx,row in runtime_data[(runtime_data.problem==1)&(runtime_data.cores>1)].iterrows():
        if has_explicit_final.loc[idx]:continue
        trials=jread(args.results.parent/'candidates'/f'{row["case"]}_p1_n{int(row["cores"])}.json',[])
        if any(t.get('method')==row['method'] and t.get('status')=='ok' and t.get('reused_singlecore') is True for t in trials):
            runtime_data.loc[idx,'total_evaluation_seconds']+=float(row.evaluation_seconds)
    runtime=[]
    for problem in (1,2,3):
        s=runtime_data[(runtime_data.problem==problem)&(runtime_data.cores>1)]
        runtime.append([problem,len(s),number(s.solver_seconds.mean(),4),number(s.solver_seconds.quantile(.95),4),number(s.total_evaluation_seconds.mean(),3),number(s.total_evaluation_seconds.quantile(.95),3)])
    timing_configs=runtime_data.loc[runtime_data.cores>1,['case','cores']].copy()
    timing_configs['certified_path']=has_explicit_final.loc[timing_configs.index]
    archived_paths=timing_configs.groupby(['case','cores']).certified_path.max()
    direct_count=int((~archived_paths).sum());certified_count=int(archived_paths.sum())
    toks['RUNTIME_TABLE']=table(['问题','多核记录数','生成均值/s','生成P95/s','总评估均值/s','总评估P95/s'],runtime,'表9 最终存档试验的候选生成与总评估耗时')+'\n总评估时间包含全部候选的评分调用；证书评分记录还包含两种场景的官方核内预构建和有效执行偏序等价检查。胜出候选若复用了经证书确认的场景B评分或单核评分，另计最终官方场景A复评时间，且不重复计入搜索耗时。旧记录按候选日志识别需要补计的整图复评，新记录使用显式最终复评字段。表中不含预计算整图单核基准、结果与证书JSON压缩读写以及框架调度的开销。所有秒数均为调用期间测得的墙钟时间，在并行进程共享CPU运行下会受资源竞争影响，不是独占CPU时间。'+f'本表混合了原路径 {direct_count} 组与证书路径 {certified_count} 组多核配置的最终存档调用（按用例与核数组合计数），不计开发验证及中断重启后丢弃的调用，也不作为两种评分引擎在相同硬件条件下的速度对比或整轮求解总耗时。\n'
    gantt_path=args.results.parent/'official_results'/f'{args.gantt_case}_p2_n5.json.gz'
    trace=jread(gantt_path);row=b[b.case==args.gantt_case]
    if trace and len(row):
        util={p:sum(o['end']-o['start'] for core in trace['per_core_timeline'] for o in core['ops'] if o['pipe']==p)/(5*trace['makespan']) for p in ('PIPE_M','PIPE_V')}
        toks['GANTT_TEXT']=f'选择 {args.gantt_case} 的场景 B 五核结果作为时间线示例。官方 Makespan 为 {integer(trace["makespan"])} cycles，相对整图单核加速比为 {number(row.speedup.iloc[0])}。按五核总可用时间计算，Cube 与 Vector 的忙碌占比分别为 {number(util["PIPE_M"]*100,2)}% 和 {number(util["PIPE_V"]*100,2)}%。两种计算 Pipe 可在同核重叠，因此上述占比不能简单相加后解释为不重叠的核心利用率。'
    else:toks['GANTT_TEXT']='【内部排版草稿：代表用例正式时间线尚未生成。】'
    toks['CONCLUSION_TEXT']=f'本文建立了覆盖三种硬件场景的切图、分核与顺序联合模型，采用结构候选构造和官方精确评估完成 100 个正式计算图的求解。场景 A 与场景 B 的 5 核平均加速比分别为 {number(a.speedup.mean())} 和 {number(b.speedup.mean())}；固定场景 B 方案后，只读 L2 的 5 核平均相对加速比为 {number(cache_ratio.mean())}，平均字节命中率为 {number(c.cache_hit_rate.mean()*100,2)}%。实验说明，图划分需要同时考虑独立分量、Pipe 负载、数据复用与容量压力；共享 Cache 的作用则需要通过同方案成对评估判断。本文结论适用于附件图集合、固定硬件参数与指定评估语义，完整方案、逐例结果和复现代码一并给出。'
    figs={
        'FIG_METHOD':('fig01_method','图2 结构感知候选构造与官方评估选优流程'),
        'FIG_DATASET':('fig02_dataset','图1 正式计算图的规模分布与计算依赖特征'),
        'FIG_SPEEDUP':('fig03_speedup','图3 场景A与场景B的平均加速比及实例四分位范围'),
        'FIG_DISTRIBUTION':('fig04_case_distribution','图4 五核加速比分布与两种场景的逐例对照'),
        'FIG_MOVEMENT':('fig06_movement','图5 五核额外数据搬运与搬运来源分解（横轴采用对称对数刻度，近零区域线性）'),
        'FIG_ABLATION':('fig09_candidate_gain','图8 最终候选组合的对照收益与方案来源'),
        'FIG_GANTT':('fig08_gantt',f'图9 {args.gantt_case}在场景B五核配置下的真实执行时间线')}
    for token,(name,cap) in figs.items():toks[token]=figure(name,cap,args)
    toks['FIG_CACHE']=figure('fig05_cache','图6 同一方案在无L2与只读Cache下的加速曲线',args)+'命中率刻画被 Cache 服务的字节比例；实际时延收益还取决于访问时序与关键路径。'+figure('fig07_cache_mechanism','图7 Cache字节命中率与时延收益的实例关系',args)
    return toks


def set_font(run,size=None,bold=None,font=BODY_FONT):
    run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    rpr=run._element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    for key,value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',font),('cs','Times New Roman')]:fonts.set(qn('w:'+key),value)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):fonts.attrib.pop(qn('w:'+key),None)


def font_style(style,size,font=BODY_FONT,bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    pr=style.element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
    for key,value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',font)]:rf.set(qn('w:'+key),value)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme'):rf.attrib.pop(qn('w:'+key),None)


def style_table(t,widths=None,appendix=False):
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[15.5/len(t.columns)]*len(t.columns)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    pr=t._tbl.tblPr; borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for old in list(borders):borders.remove(old)
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    for ri,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr();split=OxmlElement('w:cantSplit');rp.append(split)
        if ri==0:repeat=OxmlElement('w:tblHeader');rp.append(repeat)
        for cell,w in zip(row.cells,widths):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            # Pandoc's table style may leave per-cell black rules that override
            # the table-level gray grid. Remove those overrides explicitly.
            for old in list(cp.findall(qn('w:tcBorders'))):cp.remove(old)
            for side,val in [('top',45 if appendix else 70),('bottom',45 if appendix else 70),('left',65),('right',65)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            if ri==0:shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');cp.append(shade)
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0 if appendix else 1.10
                p.paragraph_format.keep_with_next=False
                for r in p.runs:set_font(r,10,bold=ri==0)


def add_numbered_equations(doc):
    paras=list(doc.paragraphs)
    for i,p in enumerate(paras):
        match=re.fullmatch(r'EQNUMBER_(\d+)',p.text.strip())
        if not match:continue
        if i==0:raise ValueError('Equation number without equation')
        prev=paras[i-1];om=prev._p.find('.//'+qn('m:oMath'))
        if om is None:raise ValueError('Pandoc did not create native OMML for equation '+match[1])
        om=deepcopy(om)
        for child in list(prev._p):
            if child.tag!=qn('w:pPr'):prev._p.remove(child)
        prev.paragraph_format.first_line_indent=Pt(0);prev.paragraph_format.tab_stops.add_tab_stop(Cm(7.6),WD_TAB_ALIGNMENT.CENTER);prev.paragraph_format.tab_stops.add_tab_stop(Cm(15.5),WD_TAB_ALIGNMENT.RIGHT)
        prev.alignment=WD_ALIGN_PARAGRAPH.LEFT;prev.add_run('\t');prev._p.append(om);r=prev.add_run('\t('+str(int(match[1]))+')');set_font(r,11)
        prev.paragraph_format.space_before=Pt(5);prev.paragraph_format.space_after=Pt(5);prev.paragraph_format.keep_together=True
        # Keep the observed short lead-in of equation (7) with its display.
        if match[1]=='7' and i>=2:
            lead=paras[i-2]
            if lead.text.strip() and not re.fullmatch(r'EQNUMBER_\d+',lead.text.strip()):
                lead.paragraph_format.keep_with_next=True
        p._p.getparent().remove(p._p)


def poststyle(path,data,args):
    doc=Document(path)
    # Pandoc's Heading names use title case; python-docx expects the built-in
    # internal names in lower case when looking up or adding headings.
    for st in doc.styles:
        if re.fullmatch(r'Heading [1-9]',st.name):
            st.element.find(qn('w:name')).set(qn('w:val'),st.name.lower())
    for section in doc.sections:
        section.page_width=Cm(21);section.page_height=Cm(29.7);section.top_margin=Cm(2.2);section.bottom_margin=Cm(2.1);section.left_margin=Cm(2.6);section.right_margin=Cm(2.6);section.header_distance=Cm(.9);section.footer_distance=Cm(1)
    for sty in ['Normal','Body Text','First Paragraph','Compact']:
        if sty in doc.styles:
            st=doc.styles[sty];font_style(st,12);pf=st.paragraph_format;pf.line_spacing=1.3;pf.space_before=Pt(0);pf.space_after=Pt(3);pf.first_line_indent=Pt(24)
    font_style(doc.styles['Title'],20,HEAD_FONT,True)
    doc.styles['Title'].paragraph_format.space_after=Pt(16)
    for k,size in [(1,15),(2,13),(3,12.5)]:
        st=doc.styles[f'Heading {k}'];font_style(st,size,HEAD_FONT,True);pf=st.paragraph_format;pf.space_before=Pt(12);pf.space_after=Pt(6);pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.keep_together=True
    for p in doc.paragraphs:
        text=p.text.strip();pf=p.paragraph_format;pf.widow_control=True
        if p.style.name=='Title':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0)
        elif p.style.name.startswith('Heading'):
            pf.first_line_indent=Pt(0)
            if text=='摘要':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if text.startswith('1 问题重述'):pf.page_break_before=True
        else:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if re.match(r'^[图表]\d+\s',text):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0);pf.line_spacing=1.1;pf.space_before=Pt(4);pf.space_after=Pt(8);pf.keep_with_next=text.startswith('表')
            for r in p.runs:set_font(r,10.5)
        if p._p.xpath('.//w:drawing'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.space_after=Pt(0);pf.space_before=Pt(6)
        if text.startswith('[') and re.match(r'^\[\d+\]',text):
            pf.first_line_indent=Pt(-18);pf.left_indent=Pt(18);pf.line_spacing=1.1
            for r in p.runs:set_font(r,10.5)
        if text.startswith('关键词'):
            pf.first_line_indent=Pt(0);pf.space_before=Pt(6)
        if text.startswith('【内部排版草稿'):
            for r in p.runs:set_font(r,11)
        # Remove pandoc theme overrides while preserving semantic bold and italic.
        for r in p.runs:
            size=r.font.size.pt if r.font.size else None
            set_font(r,size,font=HEAD_FONT if p.style.name.startswith('Heading') or p.style.name=='Title' else BODY_FONT)
    add_numbered_equations(doc)
    for shape in doc.inline_shapes:
        if shape.width>Cm(15.5):ratio=Cm(15.5)/shape.width;shape.width=Cm(15.5);shape.height=int(shape.height*ratio)
    for t in doc.tables:
        if len(t.columns)==3:widths=[3.4,9.5,2.6] if t.cell(0,0).text.strip()=='符号' else [3.2,6.2,6.1]
        elif len(t.columns)==2:widths=[6,9.5]
        elif len(t.columns)==5:widths=[3.3,3.05,3.05,3.05,3.05]
        else:widths=None
        style_table(t,widths)
    # Real appendices: landscape repeated-header official results and portrait instructions.
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Cm(29.7);sec.page_height=Cm(21);sec.left_margin=Cm(1.4);sec.right_margin=Cm(1.4);sec.top_margin=Cm(1.3);sec.bottom_margin=Cm(1.3)
    doc.add_heading('附录 A 全部用例的官方评估结果',level=1)
    note=doc.add_paragraph('每个用例列出1～5核。T为Makespan（cycles），D为相对原图新增的逻辑COPY搬运量（bytes）；A、B、C分别代表问题一、问题二、只读L2。H为Cache字节命中率。全部计数保留精确整数，H保留两位小数。');note.paragraph_format.first_line_indent=Pt(0)
    for r in note.runs:set_font(r,10)
    pivot=data.set_index(['case','cores','problem']);t=doc.add_table(rows=1,cols=9)
    for cell,txt in zip(t.rows[0].cells,['用例','核数','T_A / cycles','D_A / bytes','T_B / cycles','D_B / bytes','T_C / cycles','D_C / bytes','H / %']):cell.text=txt
    for case,n in sorted(set(zip(data.case,data.cores))):
        cells=t.add_row().cells;vals=[case,n]
        for prob in (1,2,3):
            if (case,n,prob) in pivot.index:
                row=pivot.loc[(case,n,prob)];vals.extend([integer(row.makespan),integer(row.added_copy_bytes)])
            else:vals.extend(['—','—'])
        vals.append(number(pivot.loc[(case,n,3),'cache_hit_rate']*100,2) if (case,n,3) in pivot.index else '—')
        for cell,value in zip(cells,vals):cell.text=str(value)
    style_table(t,[1.9,1.1,3.0,3.0,3.0,3.0,3.0,3.0,1.8],appendix=True)
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.PORTRAIT;sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(2.6);sec.right_margin=Cm(2.6);sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1)
    doc.add_heading('附录 B 文件说明与复现步骤',level=1)
    doc.add_paragraph('代码包包含原始数据与官方代码、调度算法、完整方案、原始评估结果、绘图数据及环境文件。所有命令均在解压后的项目根目录执行。正式统计不得修改data/config.txt或official目录内的评估脚本。')
    tt=doc.add_table(rows=1,cols=2)
    for c,s in zip(tt.rows[0].cells,['路径','用途']):c.text=s
    for name,purpose in [('data/','100个原始计算图与固定配置'),('official/','未经修改的附件官方评估代码'),('src/solver.py','结构候选、分核与子图顺序'),('src/run_certified.py','使用严格证书的候选评分与正式复评'),('src/certified_evaluate.py','有效执行偏序等价证书'),('src/run_experiments.py','原始场景评估路径（整图候选复用单核评分）'),('src/validate.py','独立方案与结果核验'),('src/analyze_data.py / plot_results.py','数据审计与论文可视化'),('src/build_paper.py','正文填数、原生公式与Word生成'),('results/plans/','1500份场景与核数对应方案'),('results/official_results/','完整官方时间线JSON压缩文件'),('results/certificates/','候选证书与等价检查记录'),('results/results.csv / candidates.csv','正式结果与候选试验明细'),('figures/','PNG、SVG、PDF和绘图CSV'),('environment.yml / requirements.txt','Conda与pip环境依赖')]:
        cells=tt.add_row().cells;cells[0].text=name;cells[1].text=purpose
    style_table(tt,[7,8.5])
    doc.add_heading('B.1 环境与运行',level=2)
    doc.add_paragraph('建议使用严格证书运行路径。只有通过偏序等价检查的候选才复用场景B评分，问题一胜出方案仍由官方场景A生成正式结果。')
    commands=['conda env create -f environment.yml','conda activate npu-scheduling','python src/run_certified.py --workers 4','python src/validate.py','python src/analyze_data.py',f'python src/plot_results.py --trace results/official_results/{args.gantt_case}_p2_n5.json.gz','python src/build_paper.py']
    for cmd in commands:
        p=doc.add_paragraph(cmd);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.1
        for r in p.runs:set_font(r,9.5)
    doc.add_paragraph('完整模拟的运行时间取决于CPU与并行进程数；已保存的同版本结果可断点续跑。若只重绘或重建论文，可直接运行最后三条命令，无需重复模拟。Word生成需要Pandoc；绘图优先使用Noto CJK或系统中文字体。')
    doc.add_paragraph('也可使用原始场景评估路径独立复现相同指标，其中整图候选复用预计算单核评分，通常耗时较长。以下命令使用单独目录，便于与证书路径的结果逐例比较：')
    for cmd in ['python src/run_experiments.py --mode full --workers 4 --output results_slow','python src/validate.py --results results_slow']:
        p=doc.add_paragraph(cmd);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.1
        for r in p.runs:set_font(r,9.5)
    doc.add_heading('B.2 单个方案复核',level=2)
    doc.add_paragraph('例如，对case_001的场景B五核方案进行独立官方复核：')
    p=doc.add_paragraph('python src/evaluate.py data/case_001.json --plan results/plans/case_001_p2_n5.json --problem 2 --output results/recheck_case_001.json');p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.1
    for r in p.runs:set_font(r,9.5)
    doc.add_paragraph('最终论文要求1500条完整记录，并通过独立校验。--allow-partial仅用于开发期内部版式检查，不用于正式输出。算法没有GPU或CUDA依赖；文中cycles是题目模拟器的执行周期，不能换算为本机Python程序的运行秒数。')
    for section in doc.sections:
        for p in section.header.paragraphs:p.text=''
        foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;foot.paragraph_format.first_line_indent=Pt(0)
        if not foot._p.xpath('.//w:fldChar'):
            run=foot.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');inst=OxmlElement('w:instrText');inst.set(qn('xml:space'),'preserve');inst.text=' PAGE ';end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.extend([begin,inst,end]);set_font(run,10)
    doc.save(path)


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--template',type=Path,default=ROOT/'paper_template.md')
    ap.add_argument('--results',type=Path,default=ROOT/'results'/'results.csv')
    ap.add_argument('--figures',type=Path,default=ROOT/'figures')
    ap.add_argument('--output',type=Path,default=ROOT.parent/'A题_多核调度完整论文.docx')
    ap.add_argument('--manuscript',type=Path,default=ROOT/'paper.md')
    ap.add_argument('--gantt-case',default='case_019')
    ap.add_argument('--allow-partial',action='store_true')
    args=ap.parse_args()
    for key in ('template','results','figures','output','manuscript'):
        setattr(args,key,getattr(args,key).resolve())
    data=pd.read_csv(args.results);validate_input(data,args)
    text=args.template.read_text(encoding='utf-8');tokens=make_tokens(data,args)
    present=set(re.findall(r'\{\{([A-Z0-9_]+)\}\}',text));missing=present-set(tokens)
    if missing:raise ValueError('Unfilled template tokens: '+str(missing))
    for key,value in tokens.items():text=text.replace('{{'+key+'}}',value)
    if re.search(r'\{\{[A-Z0-9_]+\}\}',text):raise ValueError('Unresolved template token remains')
    # TeX's legacy declaration \rm is not accepted by Pandoc's DOCX math
    # reader; its modern braced equivalent preserves the exact notation.
    text=re.sub(r'\{\\rm\s+([^{}]+)\}',lambda m:r'\mathrm{'+m.group(1)+'}',text)
    if args.allow_partial:text=text.replace('# 摘要','# 内部排版草稿\n\n本文件只用于版式检查，数据尚未完成，不得作为正式论文。\n\n# 摘要',1)
    args.manuscript.parent.mkdir(parents=True,exist_ok=True);args.manuscript.write_text(text,encoding='utf-8')
    # Pandoc silently drops LaTeX tags in DOCX. Preserve tags as markers, then
    # put native OMML and a right-aligned number in the same Word paragraph.
    def eq_repl(m):
        body=m.group(1);tag=re.search(r'\\tag\{(\d+)\}',body)
        if not tag:return m.group(0)
        return '$$'+re.sub(r'\\tag\{\d+\}','',body)+'$$\n\nEQNUMBER_'+tag.group(1)+'\n'
    tex=re.sub(r'\$\$(.*?)\$\$',eq_repl,text,flags=re.S)
    # Author to unique temporary files, never exposing half-styled DOCX to a
    # concurrent renderer or a document previewer.
    args.output.parent.mkdir(parents=True,exist_ok=True)
    tempdir=tempfile.TemporaryDirectory(prefix='npu_paper_build_',dir=args.output.parent)
    tmp=Path(tempdir.name)/'paper_pandoc.md';tmp.write_text(tex,encoding='utf-8')
    staging=Path(tempdir.name)/'paper_staging.docx'
    pandoc=shutil.which('pandoc')
    if not pandoc:raise SystemExit('Pandoc is required; install with conda install -c conda-forge pandoc')
    args.output.parent.mkdir(parents=True,exist_ok=True)
    subprocess.run([pandoc,str(tmp),'--from','markdown+tex_math_dollars','--to','docx','--standalone','--resource-path',str(ROOT),'--output',str(staging)],check=True,cwd=ROOT)
    poststyle(staging,data,args)
    from zipfile import ZipFile
    with ZipFile(staging) as z:
        xml=z.read('word/document.xml').decode('utf-8');native=xml.count('<m:oMath>')
        if 'EQNUMBER_' in xml or '{{' in xml:raise ValueError('Unresolved equation or template marker in Word')
        if native<17:raise ValueError('Native math conversion failed')
    staging.replace(args.output);tempdir.cleanup()
    print(json.dumps({'output':str(args.output),'manuscript':str(args.manuscript),'records':len(data),'native_math_objects':native,'partial':args.allow_partial},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
